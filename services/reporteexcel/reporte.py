import pandas as pd
from io import BytesIO
from datetime import datetime
from sqlalchemy import desc
from config.db import SessionLocal
from models.ticket_usuario import TicketUsuario
from models.viaje_usuario import ViajeUsuario
from models.viaticos import Viaticos
from models.monedas import Monedas
from models.viajes import Viajes
from models.conceptos import Conceptos
from models.medios_pago import MediosPago
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


async def generar_excel_tickets_service(id_viaje_usuario: int):
    """
    Reporte y resumen de tickets cargados.
    """
    
    def apply_border_to_range(ws, start_row, end_row, start_col, end_col, border_style):
        """Aplica bordes a un rango de celdas."""
        for row in range(start_row, end_row + 1):
            for col in range(start_col, end_col + 1):
                ws.cell(row=row, column=col).border = border_style
    
    db = SessionLocal()
    try:
        viaje_usuario = db.query(ViajeUsuario).filter(
            ViajeUsuario.id_viaje_usuario == id_viaje_usuario,
            ViajeUsuario.estado == True
        ).first()
        
        if not viaje_usuario:
            return {
                "success": False,
                "detail": "El viaje_usuario especificado no existe o no está activo."
            }
        
        viaje = db.query(Viajes).filter(Viajes.id_viaje == viaje_usuario.id_viaje).first()
        
        monedas = db.query(Monedas).all()
        monedas_dict = {m.id_moneda: m.nombre for m in monedas}
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Rendición"
        
        title_font = Font(bold=True, size=16, color="FFFFFF")
        title_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_label_font = Font(bold=True, size=11)
        header_fill = PatternFill(start_color="9DC3E6", end_color="9DC3E6", fill_type="solid")
        header_font = Font(bold=True, color="000000")
        section_font = Font(bold=True, size=11, color="FFFFFF")
        section_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        thick_border = Border(
            left=Side(style='medium'),
            right=Side(style='medium'),
            top=Side(style='medium'),
            bottom=Side(style='medium')
        )
        
        row = 1
        
        ws[f'A{row}'] = "PLANILLA DE RENDICIÓN DE GASTOS"
        ws[f'A{row}'].font = title_font
        ws[f'A{row}'].fill = title_fill
        ws[f'A{row}'].alignment = Alignment(horizontal='center', vertical='center')
        ws.merge_cells(f'A{row}:G{row}')
        ws.row_dimensions[row].height = 30
        
        for col in range(1, 8):
            ws.cell(row=row, column=col).border = thick_border
        
        row += 1
        
        ws[f'A{row}'] = "VIAJE/PERIODO"
        ws[f'A{row}'].font = header_label_font
        ws[f'A{row}'].border = border
        viaje_periodo = f"{viaje.nombre_viaje}" if viaje else "N/A"
        ws[f'B{row}'] = viaje_periodo
        ws[f'B{row}'].border = border
        ws.merge_cells(f'B{row}:G{row}')
        row += 1
        
        ws[f'A{row}'] = "VIAJANTE"
        ws[f'A{row}'].font = header_label_font
        ws[f'A{row}'].border = border
        ws[f'B{row}'] = f"{viaje_usuario.nombre_usuario_viajante} {viaje_usuario.apellido_usuario_viajante}" or "N/A"
        ws[f'B{row}'].border = border
        ws.merge_cells(f'B{row}:G{row}')
        row += 1
        
        ws[f'A{row}'] = "FECHA RENDICION"
        ws[f'A{row}'].font = header_label_font
        ws[f'A{row}'].border = border
        ws[f'B{row}'] = datetime.now().strftime("%Y-%m-%d")
        ws[f'B{row}'].border = border
        ws.merge_cells(f'B{row}:G{row}')
        
        for col in range(1, 8):
            for encabezado_row in range(row-2, row+1):
                ws.cell(row=encabezado_row, column=col).border = thick_border
        
        row += 2
        
        ws[f'A{row}'] = "(i) VIATICOS"
        ws[f'A{row}'].font = section_font
        ws[f'A{row}'].fill = section_fill
        ws.merge_cells(f'A{row}:G{row}')
        row += 2
        
        headers_viaticos = [
            "Efectivo entregado por tesorería",
            "Efectivo devuelto a tesorera",
            "Total de gastos rendidos",
            "Diferencia",
            "Moneda",
            "Estado"
        ]
        for col, header in enumerate(headers_viaticos, 1):
            cell = ws.cell(row=row, column=col)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        
        row += 1
        
        totales_por_moneda = {}
        for moneda in monedas:
            viaticos = db.query(Viaticos).filter(
                Viaticos.id_viaje_usuario == id_viaje_usuario,
                Viaticos.id_moneda == moneda.id_moneda
            ).all()
            
            if viaticos:
                efectivo_entregado = sum(v.monto_entregado or 0 for v in viaticos)
                efectivo_devuelto = sum(v.monto_devuelto or 0 for v in viaticos)
                
                tickets_efectivo = db.query(TicketUsuario).filter(
                    TicketUsuario.id_viaje_usuario == id_viaje_usuario,
                    TicketUsuario.id_moneda == moneda.id_moneda,
                    TicketUsuario.id_medio_pago == 1
                ).all()
                
                total_gastos_rendidos = sum(t.monto or 0 for t in tickets_efectivo)
                diferencia = efectivo_entregado - efectivo_devuelto - total_gastos_rendidos
                
                estado_viatico = "Aprobado" if diferencia == 0 else "Pendiente"
                
                ws.cell(row=row, column=1).value = efectivo_entregado
                ws.cell(row=row, column=2).value = efectivo_devuelto
                ws.cell(row=row, column=3).value = total_gastos_rendidos
                ws.cell(row=row, column=4).value = diferencia
                ws.cell(row=row, column=5).value = moneda.nombre
                ws.cell(row=row, column=6).value = estado_viatico
                
                for col in range(1, 7):
                    ws.cell(row=row, column=col).border = border
                    ws.cell(row=row, column=col).alignment = Alignment(horizontal='center', vertical='center')
                
                totales_por_moneda[moneda.id_moneda] = {
                    "nombre": moneda.nombre,
                    "entregado": efectivo_entregado,
                    "devuelto": efectivo_devuelto,
                    "gastos": total_gastos_rendidos,
                    "diferencia": diferencia
                }
                
                row += 1
        
        ws[f'A{row}'] = "Total de Gastos en Efectivo (ARS)"
        ws[f'A{row}'].font = header_label_font
        ws[f'B{row}'] = totales_por_moneda.get(1, {}).get("gastos", 0) if 1 in totales_por_moneda else 0
        row += 1
        
        ws[f'A{row}'] = "Total de Gastos en Efectivo (USD)"
        ws[f'A{row}'].font = header_label_font
        ws[f'B{row}'] = totales_por_moneda.get(2, {}).get("gastos", 0) if 2 in totales_por_moneda else 0
        row += 1
        
        ws[f'A{row}'] = "Total de Gastos en Efectivo (EUROS)"
        ws[f'A{row}'].font = header_label_font
        ws[f'B{row}'] = totales_por_moneda.get(3, {}).get("gastos", 0) if 3 in totales_por_moneda else 0
        row += 1
        
        ws[f'A{row}'] = "Total de Gastos en Efectivo (LOCAL)"
        ws[f'A{row}'].font = header_label_font
        ws[f'B{row}'] = totales_por_moneda.get(4, {}).get("gastos", 0) if 4 in totales_por_moneda else 0
        row += 2
        
        ws[f'A{row}'] = "(ii) DETALLE DE GASTOS EN EFECTIVO"
        ws[f'A{row}'].font = section_font
        ws[f'A{row}'].fill = section_fill
        ws.merge_cells(f'A{row}:G{row}')
        row += 2
        
        tickets_efectivo_all = db.query(TicketUsuario).filter(
            TicketUsuario.id_viaje_usuario == id_viaje_usuario,
            TicketUsuario.id_medio_pago == 1
        ).order_by(TicketUsuario.fecha_gasto).all()
        
        if tickets_efectivo_all:
            headers_efectivo = ["Fecha", "Comprobante", "Concepto", "Monto", "Moneda", "Estado"]
            for col, header in enumerate(headers_efectivo, 1):
                cell = ws.cell(row=row, column=col)
                cell.value = header
                cell.fill = header_fill
                cell.font = header_font
                cell.border = border
                cell.alignment = Alignment(horizontal='center', vertical='center')
            
            row += 1
            
            for ticket in tickets_efectivo_all:
                ws.cell(row=row, column=1).value = ticket.fecha_gasto
                ws.cell(row=row, column=2).value = ticket.nombre_imagen
                
                concepto = db.query(Conceptos).filter(Conceptos.id_concepto == ticket.id_concepto).first()
                concepto_nombre = concepto.nombre if concepto else ""
                ws.cell(row=row, column=3).value = concepto_nombre
                
                ws.cell(row=row, column=4).value = ticket.monto
                ws.cell(row=row, column=5).value = monedas_dict.get(ticket.id_moneda, "")
                ws.cell(row=row, column=6).value = "Aprobado" if ticket.estado == "Aprobado" else "Por aprobar"
                
                for col in range(1, 7):
                    ws.cell(row=row, column=col).border = border
                
                row += 1
                
        ws[f'A{row}'] = "Total de Gastos en ARS"
        ws[f'A{row}'].font = header_label_font
        total_ars_efectivo = sum(t.monto for t in tickets_efectivo_all if t.id_moneda == 1)
        ws[f'B{row}'] = total_ars_efectivo
        row += 1
        
        ws[f'A{row}'] = "Total de Gastos en USD"
        ws[f'A{row}'].font = header_label_font
        total_usd_efectivo = sum(t.monto for t in tickets_efectivo_all if t.id_moneda == 2)
        ws[f'B{row}'] = total_usd_efectivo
        row += 1
        
        ws[f'A{row}'] = "Total de Gastos en EUROS"
        ws[f'A{row}'].font = header_label_font
        total_eur_efectivo = sum(t.monto for t in tickets_efectivo_all if t.id_moneda == 3)
        ws[f'B{row}'] = total_eur_efectivo
        row += 1
        
        ws[f'A{row}'] = "Total de Gastos en LOCAL"
        ws[f'A{row}'].font = header_label_font
        total_local_efectivo = sum(t.monto for t in tickets_efectivo_all if t.id_moneda == 4)
        ws[f'B{row}'] = total_local_efectivo
        row += 2
        
        ws[f'A{row}'] = "(iii) DETALLE DE GASTOS EN TARJETA DE CREDITO"
        ws[f'A{row}'].font = section_font
        ws[f'A{row}'].fill = section_fill
        ws.merge_cells(f'A{row}:G{row}')
        row += 2
        
        tickets_otros = db.query(TicketUsuario).filter(
            TicketUsuario.id_viaje_usuario == id_viaje_usuario,
            TicketUsuario.id_medio_pago != 1
        ).order_by(TicketUsuario.fecha_gasto).all()
        
        if tickets_otros:
            headers_tarjeta = ["Fecha", "Comprobante", "Concepto", "Monto", "Moneda", "Estado", "Medio Pago"]
            for col, header in enumerate(headers_tarjeta, 1):
                cell = ws.cell(row=row, column=col)
                cell.value = header
                cell.fill = header_fill
                cell.font = header_font
                cell.border = border
                cell.alignment = Alignment(horizontal='center', vertical='center')
            
            row += 1
            
            for ticket in tickets_otros:
                ws.cell(row=row, column=1).value = ticket.fecha_gasto
                ws.cell(row=row, column=2).value = ticket.nombre_imagen
                
                concepto = db.query(Conceptos).filter(Conceptos.id_concepto == ticket.id_concepto).first()
                concepto_nombre = concepto.nombre if concepto else ""
                ws.cell(row=row, column=3).value = concepto_nombre
                
                ws.cell(row=row, column=4).value = ticket.monto
                ws.cell(row=row, column=5).value = monedas_dict.get(ticket.id_moneda, "")
                ws.cell(row=row, column=6).value = "Aprobado" if ticket.estado == "Aprobado" else "Por aprobar"
                
                medio_pago = db.query(MediosPago).filter(MediosPago.id_medio_pago == ticket.id_medio_pago).first()
                medio_pago_nombre = medio_pago.nombre if medio_pago else ""
                ws.cell(row=row, column=7).value = medio_pago_nombre
                
                for col in range(1, 8):
                    ws.cell(row=row, column=col).border = border
                
                row += 1
        
        ws[f'A{row}'] = "Total de Gastos en ARS"
        ws[f'A{row}'].font = header_label_font
        total_ars_tarjeta = sum(t.monto for t in tickets_otros if t.id_moneda == 1)
        ws[f'B{row}'] = total_ars_tarjeta
        row += 1
        
        ws[f'A{row}'] = "Total de Gastos en USD"
        ws[f'A{row}'].font = header_label_font
        total_usd_tarjeta = sum(t.monto for t in tickets_otros if t.id_moneda == 2)
        ws[f'B{row}'] = total_usd_tarjeta
        row += 1
        
        ws[f'A{row}'] = "Total de Gastos en EUROS"
        ws[f'A{row}'].font = header_label_font
        total_eur_tarjeta = sum(t.monto for t in tickets_otros if t.id_moneda == 3)
        ws[f'B{row}'] = total_eur_tarjeta
        row += 1
        
        ws[f'A{row}'] = "Total de Gastos en MONEDA LOCAL"
        ws[f'A{row}'].font = header_label_font
        total_local_tarjeta = sum(t.monto for t in tickets_otros if t.id_moneda == 4)
        ws[f'B{row}'] = total_local_tarjeta
        
        row += 2
        
        ws[f'A{row}'] = "FIRMA"
        ws[f'A{row}'].font = section_font
        ws[f'A{row}'].fill = section_fill
        ws.merge_cells(f'A{row}:G{row}')
        for col in range(1, 8):
            ws.cell(row=row, column=col).border = border
        
        row += 4
        
        ws[f'D{row}'] = "_" * 40
        ws[f'D{row}'].alignment = Alignment(horizontal='center', vertical='center')
        ws.merge_cells(f'D{row}:F{row}')
        
        
        row += 1
        
        ws[f'D{row}'] = "Firma del Responsable"
        ws[f'D{row}'].font = Font(bold=True, size=10)
        ws[f'D{row}'].alignment = Alignment(horizontal='center', vertical='center')
        ws.merge_cells(f'D{row}:F{row}')
        row += 3
        
        mensaje_aclaracion = "Las rendiciones en estado de 'Pendiente' no son registradas, debido a que no cuentan con los datos necesarios para ser útiles. Solicitar al viajante que las complete."
        ws[f'A{row}'] = mensaje_aclaracion
        ws[f'A{row}'].font = Font(italic=True, size=9, color="FF0000")
        ws[f'A{row}'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.merge_cells(f'A{row}:G{row}')
        ws.row_dimensions[row].height = 30
        
        aclaracion_fill = PatternFill(start_color="FFB6C1", end_color="FFB6C1", fill_type="solid")
        for col in range(1, 8):
            ws.cell(row=row, column=col).fill = aclaracion_fill
            ws.cell(row=row, column=col).border = thick_border
        
        final_row = row
        
        thick_border_all = Border(
            left=Side(style='medium'),
            right=Side(style='medium'),
            top=Side(style='medium'),
            bottom=Side(style='medium')
        )
        thick_border_left = Border(left=Side(style='medium'))
        thick_border_right = Border(right=Side(style='medium'))
        thick_border_top = Border(top=Side(style='medium'))
        thick_border_bottom = Border(bottom=Side(style='medium'))
        
        for current_row in range(1, final_row + 1):
            left_cell = ws.cell(row=current_row, column=1)
            left_border = Border(
                left=Side(style='medium'),
                top=left_cell.border.top if left_cell.border else None,
                bottom=left_cell.border.bottom if left_cell.border else None,
                right=left_cell.border.right if left_cell.border else None
            )
            left_cell.border = left_border
            
            right_cell = ws.cell(row=current_row, column=7)
            right_border = Border(
                right=Side(style='medium'),
                top=right_cell.border.top if right_cell.border else None,
                bottom=right_cell.border.bottom if right_cell.border else None,
                left=right_cell.border.left if right_cell.border else None
            )
            right_cell.border = right_border
        
        for col in range(1, 8):
            top_cell = ws.cell(row=1, column=col)
            top_border = Border(
                top=Side(style='medium'),
                left=top_cell.border.left if top_cell.border else None,
                right=top_cell.border.right if top_cell.border else None,
                bottom=top_cell.border.bottom if top_cell.border else None
            )
            top_cell.border = top_border
        
        for col in range(1, 8):
            bottom_cell = ws.cell(row=final_row, column=col)
            bottom_border = Border(
                bottom=Side(style='medium'),
                left=bottom_cell.border.left if bottom_cell.border else None,
                right=bottom_cell.border.right if bottom_cell.border else None,
                top=bottom_cell.border.top if bottom_cell.border else None
            )
            bottom_cell.border = bottom_border
        
        firma_responsable_row = final_row - 3
        border_below_firma_row = firma_responsable_row + 1
        for col in range(1, 8):
            firma_border_cell = ws.cell(row=border_below_firma_row, column=col)
            firma_border = Border(
                bottom=Side(style='medium'),
                left=firma_border_cell.border.left if firma_border_cell.border else None,
                right=firma_border_cell.border.right if firma_border_cell.border else None,
                top=firma_border_cell.border.top if firma_border_cell.border else None
            )
            firma_border_cell.border = firma_border
        
        ws.column_dimensions['A'].width = 35
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 18
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 12
        ws.column_dimensions['F'].width = 12
        ws.column_dimensions['G'].width = 15
        ws.column_dimensions['H'].width = 40
        
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        return {
            "success": True,
            "file": output
        }
    
    except Exception as e:
        return {
            "success": False,
            "detail": f"Error al generar Excel: {str(e)}"
        }
    
    finally:
        db.close()
